"""Turn raw Kobo JSON into clean tables: one row per survey (hh) and one row per family member (mem)."""
import difflib
import hashlib
import re
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

HH_COLS = [
    "_uuid", "_id", "_submission_time", "_submitted_by", "start", "end", "today", "enumerator_code", "enumerator_name",
    "enumerator_role", "interview_language", "program_region", "province", "municipality", "barangay", "hh_id",
    "informed_consent", "gps_coordinate", "hh_head_name", "hh_head_gender", "respondent_name", "respondent_relationship",
    "is_ip_family", "icc_ip_group", "hh_size_count", "hh_child_u6_count", "hh_shock_risk", "hh_challenges_top3",
    "hh_assets_owned", "housing_tenure", "land_tenure", "livelihood_sources", "primary_livelihood", "monthly_income",
    "basic_needs_sufficiency", "cultivates_farmland", "farm_hectares", "market_sold_proportion",
    "harvest_food_proportion", "harvest_shared_proportion", "severe_food_insecurity", "food_insecurity_reason",
    "produces_own_food", "home_produced_food_types", "food_from_production_months", "livestock_raised",
    "fishing_activity", "drinking_water_source", "water_access_status", "water_collection_time",
    "water_safety_confidence", "toilet_quality", "family_nutrition_ability", "family_voice",
    "leaders_assistance_confidence", "climate_resilience_confidence", "ipra_awareness", "school_community_engagement",
    "sustainable_practices", "environmental_practices", "community_org_participation",
    "governance_cycle_participation", "school_engagement_type_answer", "manapack_consistency",
    "caregiver_feeding_practice", "family_development_score", "hds_band",
]
MEM_COLS = [
    "_uuid", "member_index", "member_full_name", "member_relationship", "member_sex", "member_dob", "member_age",
    "member_nutrition_status", "member_plw_status", "member_civil_status", "member_same_icc", "member_icc_specify",
    "member_schooling_status", "member_educ_attainment", "attending_regularly", "dropout_risk", "has_scholarship",
    "scholarship_provider", "member_occupation", "member_occupation_other", "member_employment_status",
    "member_special_status", "member_disability_type", "member_leadership_role",
]


def _strip(key):
    return key.split("/")[-1] if isinstance(key, str) else key


def normalize_record(rec):
    """Remove group prefixes and normalize unique identifiers."""
    out = {}
    for k, v in rec.items():
        if k.startswith("_"):
            out[k] = v
            continue
        nk = _strip(k)
        if nk == "uuid" or k == "formhub/uuid":
            out["_uuid"] = v
            continue
        if isinstance(v, list) and v and isinstance(v[0], dict):
            out[nk] = [{_strip(kk): vv for kk, vv in item.items()} for item in v]
        else:
            out[nk] = v

    # --- ADD MUNICIPALITY NORMALIZATION HERE ---
    if "municipality" in out and isinstance(out["municipality"], str):
        out["municipality"] = out["municipality"].strip().title()

    return out


def load_choice_labels(xlsx_path):
    """{list_name: {code: English label}} from an XLSForm."""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb["choices"]
    rows = list(ws.iter_rows(values_only=True))
    head = list(rows[0])
    li, ni = head.index("list_name"), head.index("name")
    label_col = next(i for i, h in enumerate(head) if h and str(h).startswith("label"))
    out = {}
    for r in rows[1:]:
        if r[li] and r[ni] is not None:
            out.setdefault(r[li], {})[str(r[ni])] = r[label_col] or str(r[ni])
    return out


def load_item_labels(xlsx_path):
    """{field_name: English label} from the survey sheet of an XLSForm (used for monitoring items)."""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    rows = list(wb["survey"].iter_rows(values_only=True))
    head = list(rows[0])
    ni = head.index("name")
    li = next(i for i, h in enumerate(head) if h and str(h).startswith("label"))
    def tidy(x):
        x = str(x).replace("**", "")
        x = x.replace("${member_full_name}", "[family member]").replace("${hh_head_name}", "[family head]")
        return re.sub(r"\$\{(\w+)\}", r"[\1]", x).strip()
    return {r[ni]: tidy(r[li]) for r in rows[1:] if r[ni] and r[li]}


def load_directory(path):
    """Enumerator / CF / monitor directory -> DataFrame(code, name, title, area, barangay)."""
    df = pd.read_excel(path, dtype=str).fillna("")
    df.columns = [c.split(" (")[0].strip().lower() for c in df.columns]
    df = df.rename(columns={"role": "title", "area": "home_area"})
    for c in ("code", "name", "title", "home_area", "barangay"):
        if c not in df.columns:
            df[c] = ""
    df["label"] = np.where(df["name"] != "", df["code"] + " - " + df["name"], df["code"])
    return df[["code", "name", "title", "home_area", "barangay", "label"]]


def _ensure(df, cols):
    for c in cols:
        if c not in df.columns:
            df[c] = np.nan
    return df


def _parse_gps(s):
    try:
        p = str(s).split()
        return float(p[0]), float(p[1]), float(p[3]) if len(p) > 3 else np.nan
    except Exception:
        return np.nan, np.nan, np.nan



# ---------------------------------------------------------------------------------------------
# Enumerator resolution.  Some FDAF versions (e.g. the MINSOC forms) have no enumerator CODE - the
# enumerator typed their name, spelled differently each time ("Tadi,Analiza Montilla",
# "Tadi Analiza Montilla", "ROSEJANE LASIB MANDULAN").  We match typed names to the directory, and
# group the rest by a normalised spelling so that each person still appears once.
# ---------------------------------------------------------------------------------------------
def _tokens(name):
    """Lower-case word tokens of a name, without punctuation and one-letter initials."""
    t = re.sub(r"[^\w\s]", " ", str(name).lower()).split()
    return [w for w in t if len(w) > 1]


def _pretty_name(name):
    return re.sub(r"\s+", " ", re.sub(r"[,;]", " ", str(name))).strip().title()


def _match_directory(name, dir_tokens):
    """Return the directory code whose name words are ALL contained in the typed name (unique match only)."""
    toks = set(_tokens(name))
    if len(toks) < 2:
        return ""
    hits = [code for code, dt in dir_tokens.items() if len(dt) >= 2 and set(dt) <= toks]
    if len(hits) == 1:
        return hits[0]
    if hits:
        return ""
    # typo tolerance: every directory word must have a near-identical word in the typed name
    def near(w):
        return bool(difflib.get_close_matches(w, toks, n=1, cutoff=0.85))
    fuzzy = [code for code, dt in dir_tokens.items() if len(dt) >= 2 and all(near(w) for w in dt)]
    return fuzzy[0] if len(fuzzy) == 1 else ""


def resolve_enumerators(hh, directory):
    """Fill enumerator_code / enumerator_name for rows that only have a typed name."""
    if "enumerator_name" not in hh.columns:
        hh["enumerator_name"] = np.nan
    code = hh["enumerator_code"].fillna("").astype(str).str.strip()
    # codes picked from a list look like "E039" or "E039 - Name" depending on the export; keep the code part
    code = code.str.extract(r"^\s*((?:E|CF)\d+)", expand=False).fillna(code.where(~code.str.contains(r"\s-\s"), ""))
    name = hh["enumerator_name"].fillna("").astype(str).str.strip()
    hh["enumerator_name_typed"] = name

    d_tokens = {r.code: _tokens(r.name) for r in directory.itertuples() if re.match(r"^(E|CF)\d+$", str(r.code))}
    cache, auto_labels = {}, {}
    for i in hh.index[code.eq("") & name.ne("")]:
        nm = name[i]
        key = " ".join(sorted(_tokens(nm))) or nm.lower()
        if key not in cache:
            matched = _match_directory(nm, d_tokens)
            if matched:
                cache[key] = matched
            else:  # not in the directory: stable synthetic code so the person is still tracked
                cache[key] = "X-" + hashlib.md5(key.encode("utf-8")).hexdigest()[:5].upper()
                auto_labels.setdefault(cache[key], {}).setdefault(_pretty_name(nm), 0)
        code[i] = cache[key]
        if cache[key] in auto_labels:
            auto_labels[cache[key]][_pretty_name(nm)] = auto_labels[cache[key]].get(_pretty_name(nm), 0) + 1
    hh["enumerator_code"] = code
    hh["_auto_enum_labels"] = hh["enumerator_code"].map({k: max(v, key=v.get) for k, v in auto_labels.items()})
    return hh


def build_fdaf_tables(records, cfg, labels, directory):
    tz = cfg.get("timezone", "Asia/Manila")
    hh_rows, mem_rows = [], []
    for rec in records:
        r = normalize_record(rec)
        members = r.pop("hh_members", None) or []
        r["n_members"] = len(members)
        hh_rows.append(r)
        for i, m in enumerate(members, 1):
            m = dict(m)
            m["_uuid"] = r.get("_uuid")
            m["member_index"] = i
            mem_rows.append(m)
    
    hh_df = pd.DataFrame(hh_rows)
    hh_df = hh_df.loc[:, ~hh_df.columns.duplicated()]
    hh = _ensure(hh_df, HH_COLS + ["n_members", "_source_form"])

    mem_df = pd.DataFrame(mem_rows)
    mem_df = mem_df.loc[:, ~mem_df.columns.duplicated()]
    mem = _ensure(mem_df, MEM_COLS)

    # column names starting with "_" break DataFrame.itertuples(), so use plain names
    ren = {"_uuid": "uuid", "_id": "kobo_id", "_submission_time": "submission_time", "_submitted_by": "submitted_by", "_source_form": "source_form"}
    hh = hh.rename(columns=ren)
    mem = mem.rename(columns=ren)

    # Clean any duplicate column names created after renaming
    hh = hh.loc[:, ~hh.columns.duplicated()]
    mem = mem.loc[:, ~mem.columns.duplicated()]

    for c in ("start", "end", "submission_time"):
        hh[c + "_dt"] = pd.to_datetime(hh[c], errors="coerce", utc=True).dt.tz_convert(tz)
    today = pd.to_datetime(hh["today"], errors="coerce")
    fallback = hh["start_dt"].dt.tz_localize(None).dt.normalize()
    sub_day = hh["submission_time_dt"].dt.tz_localize(None).dt.normalize()
    hh["interview_date"] = today.fillna(fallback).fillna(sub_day).dt.normalize()
    hh["duration_min"] = (hh["end_dt"] - hh["start_dt"]).dt.total_seconds() / 60
    gps = hh["gps_coordinate"].apply(_parse_gps)
    hh["lat"] = [g[0] for g in gps]
    hh["lon"] = [g[1] for g in gps]
    hh["gps_acc"] = [g[2] for g in gps]
    for c in ("monthly_income", "farm_hectares", "hh_size_count", "hh_child_u6_count", "family_development_score"):
        hh[c] = pd.to_numeric(hh[c], errors="coerce")

    # labels for places
    hh["region_label"] = hh["program_region"].map(labels.get("program_region", {})).fillna(hh["program_region"])
    hh["province_label"] = hh["province"].map(labels.get("program_province", {})).fillna(hh["province"])
    hh["municipality_label"] = hh["municipality"].map(labels.get("program_municipality", {})).fillna(hh["municipality"])
    hh["barangay_label"] = hh["barangay"].map(labels.get("program_barangay", {})).fillna(hh["barangay"])

    # enumerator / role from directory
    d = directory.set_index("code")
    hh["enumerator_code"] = hh["enumerator_code"].fillna("").astype(str)
    hh = resolve_enumerators(hh, directory)
    hh["enumerator_label"] = (hh["enumerator_code"].map(d["label"]).fillna(hh["_auto_enum_labels"])
                              .fillna(hh["enumerator_code"]))
    hh["enumerator_label"] = hh["enumerator_label"].replace("", "Unknown enumerator")
    hh["title"] = hh["enumerator_code"].map(d["title"])
    hh["title"] = hh["title"].fillna(hh["enumerator_role"]).fillna("Unknown").replace("", "Unknown")
    hh["home_area"] = hh["enumerator_code"].map(d["home_area"]).fillna("")
    hh["consented"] = hh["informed_consent"].astype(str).eq("1")

    # member table: numeric age, dob, parent survey fields
    mem["member_age"] = pd.to_numeric(mem["member_age"], errors="coerce")
    mem["dob"] = pd.to_datetime(mem["member_dob"], errors="coerce")
    keep = hh[["uuid", "hh_id", "enumerator_code", "interview_date", "municipality_label", "barangay_label"]]
    # Ensure columns are unique before merging to prevent duplicate label errors
    mem = mem.loc[:, ~mem.columns.duplicated()]
    keep = keep.loc[:, ~keep.columns.duplicated()]
    
    mem = mem.merge(keep, on="uuid", how="left")
    mem["age_calc"] = np.floor((mem["interview_date"] - mem["dob"]).dt.days / 365.25)
    return hh.reset_index(drop=True), mem.reset_index(drop=True)
